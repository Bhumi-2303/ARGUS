from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document('ARGUS_paper_v2_representation_resolution.docx')

doc.add_page_break()
heading = doc.add_paragraph()
heading_run = heading.add_run('Supplementary Figures: Verified Audit Results')
heading_run.bold = True
heading_run.font.size = Pt(16)
heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

figures_to_add = [
    {
        'path': 'scratch_nr04/figures/NR04_ROC.png',
        'caption': 'Figure S1: Native SCADA (D3) Receiver Operating Characteristic (ROC) curve.'
    },
    {
        'path': 'scratch_rep01/Volumes/BLACK-BOX/ARGUS/experiment_execution/neural_robustness/representation_extension/figures/REP01_ROC_comparison.png',
        'caption': 'Figure S2: REP-01 Representation Extension ROC comparison (evaluating native flow, protocol-aware, and temporal representations).'
    },
    {
        'path': 'scratch_nr04/figures/NR04_SHAP_summary.png',
        'caption': 'Figure S3: SHAP Summary for Native SCADA representation, highlighting the top driving features.'
    },
    {
        'path': 'scratch_rep01/Volumes/BLACK-BOX/ARGUS/experiment_execution/neural_robustness/representation_extension/figures/REP01_SHAP_summary.png',
        'caption': 'Figure S4: SHAP Summary for REP-01 Protocol-Aware representation, showing attribution of protocol and temporal features.'
    },
    {
        'path': 'scratch_nr04/figures/NR04_model_capacity_comparison.png',
        'caption': 'Figure S5: Model capacity comparison on the Native SCADA representation.'
    }
]

for fig in figures_to_add:
    try:
        doc.add_paragraph() # spacing
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(fig['path'], width=Inches(5.5))
        
        cap_para = doc.add_paragraph()
        cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = cap_para.add_run(fig['caption'])
        cap_run.italic = True
    except Exception as e:
        print(f"Error adding {fig['path']}: {e}")

doc.save('ARGUS_paper_v2_representation_resolution_with_images.docx')
