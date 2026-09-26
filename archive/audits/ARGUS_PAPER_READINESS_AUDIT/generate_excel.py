#!/usr/bin/env python3
"""Generate FINAL_PAPER_READINESS_AUDIT.xlsx from all CSV audit files."""

import csv
import os

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    print("openpyxl not available. Skipping Excel generation.")

def csv_to_sheet(wb, sheet_name, csv_path):
    """Read a CSV file and write it to a named worksheet."""
    ws = wb.create_sheet(title=sheet_name[:31])  # Excel max sheet name length
    
    with open(csv_path, 'r', encoding='utf-8') as f:
        reader = csv.reader(f)
        for row_idx, row in enumerate(reader, 1):
            for col_idx, cell_value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=cell_value)
                if row_idx == 1:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                    cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(wrap_text=True, vertical='top')
    
    # Auto-width columns
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, min(len(str(cell.value)), 60))
        ws.column_dimensions[col_letter].width = max(max_length + 2, 12)
    
    return ws

def main():
    if not HAS_OPENPYXL:
        print("Cannot generate Excel without openpyxl. Install with: pip install openpyxl")
        return
    
    audit_dir = "/Volumes/BLACK-BOX/ARGUS/ARGUS_PAPER_READINESS_AUDIT"
    output_path = os.path.join(audit_dir, "FINAL_PAPER_READINESS_AUDIT.xlsx")
    
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    csv_files = [
        ("Experiment Inventory", "experiment_inventory.csv"),
        ("Evidence Matrix", "evidence_matrix.csv"),
        ("Claim Audit", "claim_audit.csv"),
        ("Metric Audit", "metric_audit.csv"),
        ("Statistical Audit", "statistical_audit.csv"),
        ("Missing Experiments", "missing_experiments.csv"),
        ("Table Plan", "table_plan.csv"),
        ("Figure Plan", "figure_plan.csv"),
        ("Section Readiness", "section_readiness.csv"),
    ]
    
    for sheet_name, csv_filename in csv_files:
        csv_path = os.path.join(audit_dir, csv_filename)
        if os.path.exists(csv_path):
            csv_to_sheet(wb, sheet_name, csv_path)
            print(f"  Added sheet: {sheet_name}")
        else:
            print(f"  WARNING: {csv_filename} not found, skipping")
    
    wb.save(output_path)
    print(f"\nExcel workbook saved to: {output_path}")

if __name__ == "__main__":
    main()
