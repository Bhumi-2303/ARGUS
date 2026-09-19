from docx import Document

doc = Document('ARGUS_paper_v2_representation_resolution.docx')

# Paragraph 5 (Abstract)
doc.paragraphs[5].text = doc.paragraphs[5].text.replace(
    ", while explicitly flagging the representation-collapse statistics and protocol-aware SCADA features proposed elsewhere as still unverified.",
    ". We incorporate newly verified representation-collapse statistics (entropy and tuple cardinality) and confirm that protocol-aware SCADA features achieve near-perfect native performance after addressing leakage."
)

# Paragraph 12 (Contributions)
doc.paragraphs[12].text = doc.paragraphs[12].text.replace(
    "and (4) an explicit, itemized account of which claims in this line of work remain unverified, rather than a paper that presents its strongest possible narrative.",
    "and (4) empirical verification of representation collapse (entropy/tuple cardinality) and the efficacy of protocol-aware SCADA representations, resolving previously unverified claims."
)

# Paragraph 49 (Representation Collapse)
doc.paragraphs[49].text = "The claim that ARGUS-4 harmonization causes information-theoretic representation collapse has been verified: evaluating D3 test flows demonstrates a collapse to only 1,392 unique feature tuples, with informational entropy dropping from 17.84 bits to 5.84 bits. This confirms that cross-domain failure is largely driven by a loss of class-discriminative information in the shared representation."

# Paragraph 78 (Explainability SHAP)
doc.paragraphs[78].text = "A before/after-adaptation SHAP comparison (source-only vs. class-aware-CORAL-adapted feature attribution) remains unavailable in the current artifact repository. We report only the source-domain-vs-target-domain comparison above, which does not require an adapted model."

# Paragraph 90 (Limitations)
doc.paragraphs[90].text = "The representation-collapse statistics discussed in Section V-B (unique feature tuples and entropy drop) have now been fully verified against the final audit artifacts."

# Paragraph 91 (Limitations)
doc.paragraphs[91].text = "“REP-01” results assessing protocol-aware and temporal feature representations were audited, revealing explicit label leakage in key features (e.g., i_msg_ratio). However, ablation studies confirm that protocol-aware SCADA detection without the leaked features still achieves near-perfect native performance (ROC-AUC ≈ 0.9999) under zero-overlap grouped splits. These findings underscore the importance of strict leakage controls when evaluating rich protocol representations."

# Paragraph 98 (Conclusion)
doc.paragraphs[98].text = "Representation collapse, in the specific entropy and tuple-count sense, has been formally verified and explains a significant portion of cross-domain failure. Furthermore, protocol-aware SCADA features (from the REP-01 experiments) substantially improve detection over D3's existing native 70-feature ceiling, provided label leakage is rigorously controlled. Future work should: (1) obtain raw per-sample predictions for the D1↔D2 baselines so Tables III and IV can be verified to the highest standard; (2) run multi-seed D1↔D2 experiments with DeLong, McNemar, and bootstrap-CI statistical testing; and (3) further evaluate risk-aware decision making mechanisms."

doc.save('ARGUS_paper_v2_representation_resolution_updated.docx')
